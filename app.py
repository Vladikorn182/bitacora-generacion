import io
import os
import re
import hashlib
import unicodedata
from datetime import datetime

import pandas as pd
import streamlit as st

# ============================================================
# CONFIGURACIÓN
# ============================================================

st.set_page_config(
    page_title="Seguimiento de Objetivos de Socios",
    page_icon="🎯",
    layout="wide",
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

SOCIOS_FILE = os.path.join(BASE_DIR, "socios_base.csv")
ACUMULADOS_FILE = os.path.join(BASE_DIR, "acumulados.csv")
AUDITORIA_FILE = os.path.join(BASE_DIR, "auditoria_cargas.csv")

FALLBACK_SOCIOS = [
    {"POS_CODE": "59509", "POS_OWNER": "ADRIANA PAOLA VILLAFUERTE GUERRA", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 90},
    {"POS_CODE": "89859", "POS_OWNER": "JOSE PABLO FERNANDEZ", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 30},
    {"POS_CODE": "88463", "POS_OWNER": "ALICIA GRACIELA ZAMORA BUEZO", "CATEGORIA": "PYME DIGITAL", "OBJETIVO": 30},
    {"POS_CODE": "78340", "POS_OWNER": "OLIVIA SANCHEZ QUISPE", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 30},
    {"POS_CODE": "83457", "POS_OWNER": "ESTRELLA BELEN QUISPE FLORES", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 30},
    {"POS_CODE": "89326", "POS_OWNER": "PALMIRA SELAES", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 20},
    {"POS_CODE": "72210", "POS_OWNER": "GUADALUPE APAZA VILA", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 20},
    {"POS_CODE": "86737", "POS_OWNER": "ANAHI OINCA", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 15},
    {"POS_CODE": "78349", "POS_OWNER": "FRANZ REYNALDO TORREZ CALLE", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 15},
    {"POS_CODE": "91208", "POS_OWNER": "DANNY QUISBERT MENDOZA", "CATEGORIA": "PYME DIGITAL", "OBJETIVO": 5},
    {"POS_CODE": "91283", "POS_OWNER": "SOLAGEL QUENTA GUTIERREZ", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 10},
    {"POS_CODE": "91262", "POS_OWNER": "WARNES RIVERA CHUQUIMIA", "CATEGORIA": "SOCIO HOGAR", "OBJETIVO": 10},
    {"POS_CODE": "91207", "POS_OWNER": "GUALBERTO FERNANDO SANJINES", "CATEGORIA": "PYME DIGITAL", "OBJETIVO": 10},
    {"POS_CODE": "88874", "POS_OWNER": "GUSTAVO CALLEJAS", "CATEGORIA": "PYME DIGITAL", "OBJETIVO": 10},
]


# ============================================================
# UTILIDADES
# ============================================================

def normalize_text(value):
    """Normaliza texto para comparar nombres sin depender de mayúsculas,
    tildes, espacios o signos."""
    if value is None:
        return ""
    text = str(value).strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_code(value):
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    text = re.sub(r"\D", "", text)
    return text


def safe_amount(value):
    if value is None or pd.isna(value):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    text = text.replace(" ", "")

    # Maneja 1.234,50 y 1234.50
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")

    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else 0.0


def file_hash(file_bytes):
    return hashlib.sha256(file_bytes).hexdigest()


def initialize_files():
    if not os.path.exists(SOCIOS_FILE):
        pd.DataFrame(FALLBACK_SOCIOS).to_csv(
            SOCIOS_FILE, index=False, encoding="utf-8-sig"
        )

    if not os.path.exists(ACUMULADOS_FILE):
        pd.DataFrame(
            columns=["POS_CODE", "GENERADO"]
        ).to_csv(
            ACUMULADOS_FILE, index=False, encoding="utf-8-sig"
        )

    if not os.path.exists(AUDITORIA_FILE):
        pd.DataFrame(
            columns=[
                "fecha_hora",
                "usuario",
                "archivo",
                "tipo",
                "hash_archivo",
                "pos_code",
                "socio",
                "monto_sumado",
                "estado",
            ]
        ).to_csv(
            AUDITORIA_FILE, index=False, encoding="utf-8-sig"
        )


def load_socios():
    initialize_files()
    df = pd.read_csv(SOCIOS_FILE, dtype=str, keep_default_na=False)

    required = ["POS_CODE", "POS_OWNER", "CATEGORIA", "OBJETIVO"]
    for col in required:
        if col not in df.columns:
            df[col] = ""

    df["POS_CODE"] = df["POS_CODE"].apply(normalize_code)
    df["POS_OWNER"] = df["POS_OWNER"].astype(str).str.strip()
    df["CATEGORIA"] = df["CATEGORIA"].astype(str).str.strip()
    df["OBJETIVO"] = df["OBJETIVO"].apply(safe_amount)

    df = df[df["POS_CODE"] != ""].copy()
    df = df.drop_duplicates(subset=["POS_CODE"], keep="last")
    return df.reset_index(drop=True)


def load_acumulados():
    initialize_files()
    df = pd.read_csv(
        ACUMULADOS_FILE,
        dtype={"POS_CODE": str},
        keep_default_na=False,
    )

    if "POS_CODE" not in df.columns:
        df["POS_CODE"] = ""
    if "GENERADO" not in df.columns:
        df["GENERADO"] = 0

    df["POS_CODE"] = df["POS_CODE"].apply(normalize_code)
    df["GENERADO"] = df["GENERADO"].apply(safe_amount)

    return df.groupby("POS_CODE", as_index=False)["GENERADO"].sum()


def load_auditoria():
    initialize_files()
    return pd.read_csv(
        AUDITORIA_FILE,
        dtype=str,
        keep_default_na=False,
    )


def save_acumulados(df):
    out = df.copy()
    out["POS_CODE"] = out["POS_CODE"].apply(normalize_code)
    out["GENERADO"] = out["GENERADO"].apply(safe_amount)
    out = out.groupby("POS_CODE", as_index=False)["GENERADO"].sum()
    out.to_csv(
        ACUMULADOS_FILE,
        index=False,
        encoding="utf-8-sig",
    )


def append_auditoria(rows):
    current = load_auditoria()
    addition = pd.DataFrame(rows)
    result = pd.concat([current, addition], ignore_index=True)
    result.to_csv(
        AUDITORIA_FILE,
        index=False,
        encoding="utf-8-sig",
    )


def save_objectives_from_excel(uploaded_bytes):
    try:
        df = pd.read_excel(
            io.BytesIO(uploaded_bytes),
            sheet_name="Objetivos",
            dtype=str,
        )
    except Exception as exc:
        raise ValueError(
            f"No pude leer la hoja 'Objetivos': {exc}"
        )

    normalized_columns = {
        normalize_text(col): col
        for col in df.columns
    }

    def find_col(*candidates):
        for candidate in candidates:
            key = normalize_text(candidate)
            if key in normalized_columns:
                return normalized_columns[key]
        return None

    pos_col = find_col("POS_CODE", "POS CODE", "EH", "POS")
    owner_col = find_col("POS_OWNER", "POS OWNER", "SOCIO", "NOMBRE")
    cat_col = find_col("CATEGORIA", "CATEGORÍA", "CATEGORIA SOCIO")
    obj_col = find_col("BU JUNIO", "OBJETIVO", "META")

    missing = []
    if not pos_col:
        missing.append("POS_CODE")
    if not owner_col:
        missing.append("POS_OWNER")
    if not cat_col:
        missing.append("CATEGORIA")
    if not obj_col:
        missing.append("BU JUNIO / OBJETIVO")

    if missing:
        raise ValueError(
            "Faltan columnas requeridas: " + ", ".join(missing)
        )

    out = pd.DataFrame(
        {
            "POS_CODE": df[pos_col].apply(normalize_code),
            "POS_OWNER": df[owner_col].astype(str).str.strip(),
            "CATEGORIA": df[cat_col].astype(str).str.strip(),
            "OBJETIVO": df[obj_col].apply(safe_amount),
        }
    )

    out = out[out["POS_CODE"] != ""].copy()
    out = out.drop_duplicates("POS_CODE", keep="last")

    if out.empty:
        raise ValueError("El Excel no contiene socios válidos.")

    out.to_csv(
        SOCIOS_FILE,
        index=False,
        encoding="utf-8-sig",
    )

    return out


# ============================================================
# VÍA A — LECTOR DE ARCHIVOS DE VENTAS
# ============================================================

def detect_sales_columns(df):
    normalized = {
        normalize_text(col): col
        for col in df.columns
    }

    pos_candidates = [
        "POS_CODE",
        "POS CODE",
        "POSCODE",
        "EH",
        "EH CODE",
        "CODIGO",
        "CODIGO POS",
        "POS",
    ]

    amount_candidates = [
        "MONTO",
        "GENERADO",
        "VENTAS",
        "VENTA",
        "GROSSADD",
        "GROSS ADD",
        "CANTIDAD",
        "TOTAL",
        "BU",
    ]

    pos_col = None
    amount_col = None

    for candidate in pos_candidates:
        key = normalize_text(candidate)
        if key in normalized:
            pos_col = normalized[key]
            break

    for candidate in amount_candidates:
        key = normalize_text(candidate)
        if key in normalized:
            amount_col = normalized[key]
            break

    # Si no encontró monto por nombre, busca una columna numérica.
    if amount_col is None:
        numeric_candidates = []
        for col in df.columns:
            converted = pd.to_numeric(
                df[col].astype(str).str.replace(",", ".", regex=False),
                errors="coerce",
            )
            if converted.notna().sum() >= max(1, len(df) * 0.5):
                numeric_candidates.append(col)

        if len(numeric_candidates) == 1:
            amount_col = numeric_candidates[0]

    return pos_col, amount_col


def read_sales_file(file_bytes, filename):
    lower = filename.lower()

    try:
        if lower.endswith(".csv"):
            try:
                df = pd.read_csv(
                    io.BytesIO(file_bytes),
                    dtype=str,
                    encoding="utf-8-sig",
                )
            except UnicodeDecodeError:
                df = pd.read_csv(
                    io.BytesIO(file_bytes),
                    dtype=str,
                    encoding="latin-1",
                )
        elif lower.endswith((".xlsx", ".xls")):
            df = pd.read_excel(
                io.BytesIO(file_bytes),
                dtype=str,
            )
        else:
            raise ValueError(
                "Formato no soportado. Usa CSV o Excel."
            )
    except Exception as exc:
        raise ValueError(
            f"No pude leer el archivo de ventas: {exc}"
        )

    if df.empty:
        raise ValueError("El archivo de ventas está vacío.")

    pos_col, amount_col = detect_sales_columns(df)

    if not pos_col:
        raise ValueError(
            "No encontré una columna POS_CODE o EH."
        )

    if not amount_col:
        raise ValueError(
            "No encontré una columna de monto generado."
        )

    out = pd.DataFrame(
        {
            "identificador": df[pos_col].apply(normalize_code),
            "monto": df[amount_col].apply(safe_amount),
        }
    )

    out["fuente"] = "archivo"

    out = out[
        (out["identificador"] != "")
        & (out["monto"] != 0)
    ].copy()

    return out


# ============================================================
# VÍA B — OCR
# ============================================================

def ocr_extract_rows(image_bytes):
    """
    Punto aislado de OCR.

    Si en el futuro se cambia pytesseract por otro motor,
    solo debe sustituirse esta función.

    Devuelve filas con:
      identificador
      nombre_ocr
      monto
      fuente
      observacion
    """
    try:
        import pytesseract
        from PIL import Image, ImageEnhance, ImageOps
    except Exception:
        return [], (
            "pytesseract/Pillow no está disponible. "
            "Puedes seguir usando la Vía A o introducir las filas manualmente."
        )

    try:
        image = Image.open(
            io.BytesIO(image_bytes)
        ).convert("RGB")

        # Preprocesamiento simple para capturas de tablas.
        image = ImageOps.grayscale(image)
        image = ImageEnhance.Contrast(image).enhance(2.0)

        text = pytesseract.image_to_string(
            image,
            config="--psm 6",
            lang="spa+eng",
        )
    except Exception as exc:
        return [], f"OCR falló: {exc}"

    rows = []

    for line in text.splitlines():
        line = line.strip()

        if not line:
            continue

        if re.search(
            r"\b(TOTAL|EJECUTIVO|SOCIO|MTD|D[IÍ]A|RANKING|POS_OWNER)\b",
            line,
            re.IGNORECASE,
        ):
            continue

        numbers = re.findall(
            r"(?<!\d)(\d{1,8}(?:[.,]\d+)?)",
            line,
        )

        if not numbers:
            continue

        # Prioridad: código POS de 4-8 dígitos.
        code = ""
        for n in numbers:
            candidate = normalize_code(n)
            if 4 <= len(candidate) <= 8:
                code = candidate
                break

        # Para capturas tabulares, el último número suele ser
        # la cifra que interesa. Si hay código y varias columnas,
        # tomamos el último número.
        amount = safe_amount(numbers[-1])

        # Texto no numérico para poder corregir / identificar por nombre.
        text_part = re.sub(
            r"^\s*\d+\s*[.)-]?\s*",
            "",
            line,
        )
        text_part = re.sub(
            r"\d+(?:[.,]\d+)?",
            " ",
            text_part,
        )
        text_part = re.sub(
            r"\s+",
            " ",
            text_part,
        ).strip(" -|:")

        if code or text_part:
            rows.append(
                {
                    "identificador": code,
                    "nombre_ocr": text_part,
                    "monto": amount,
                    "fuente": "ocr",
                    "observacion": "",
                }
            )

    if not rows:
        return [], (
            "El OCR no pudo identificar filas. "
            "Puedes introducirlas manualmente."
        )

    return rows, text


# ============================================================
# CRUCE Y PREVISUALIZACIÓN
# ============================================================

def build_lookup(socios):
    by_code = {}
    by_name = {}

    for _, row in socios.iterrows():
        code = normalize_code(row["POS_CODE"])
        name = normalize_text(row["POS_OWNER"])

        if code:
            by_code[code] = row

        if name:
            by_name[name] = row

    return by_code, by_name


def resolve_rows(rows, socios):
    by_code, by_name = build_lookup(socios)

    resolved = []
    unidentified = []

    for row in rows:
        code = normalize_code(row.get("identificador", ""))
        name_ocr = str(row.get("nombre_ocr", "")).strip()
        amount = safe_amount(row.get("monto", 0))
        source = row.get("fuente", "manual")
        observation = row.get("observacion", "")

        matched = None
        match_type = ""

        if code and code in by_code:
            matched = by_code[code]
            match_type = "POS_CODE"

        elif name_ocr:
            normalized_name = normalize_text(name_ocr)

            if normalized_name in by_name:
                matched = by_name[normalized_name]
                match_type = "NOMBRE"

            else:
                # Coincidencia simple por inclusión para OCR imperfecto.
                possible = [
                    (key, value)
                    for key, value in by_name.items()
                    if normalized_name
                    and (
                        normalized_name in key
                        or key in normalized_name
                    )
                ]

                if len(possible) == 1:
                    matched = possible[0][1]
                    match_type = "NOMBRE_PARCIAL"

        if matched is None:
            unidentified.append(
                {
                    "identificador": code,
                    "nombre_ocr": name_ocr,
                    "monto": amount,
                    "fuente": source,
                    "observacion": (
                        observation
                        or "No identificado"
                    ),
                }
            )
            continue

        resolved.append(
            {
                "POS_CODE": normalize_code(matched["POS_CODE"]),
                "socio": matched["POS_OWNER"],
                "categoria": matched["CATEGORIA"],
                "monto": amount,
                "fuente": source,
                "match": match_type,
                "observacion": observation,
            }
        )

    return pd.DataFrame(resolved), pd.DataFrame(unidentified)


def build_preview_table(resolved, socios):
    if resolved.empty:
        return resolved

    grouped = (
        resolved.groupby(
            ["POS_CODE", "socio", "categoria"],
            as_index=False,
        )["monto"]
        .sum()
        .rename(columns={"monto": "AGREGAR"})
    )

    return grouped


# ============================================================
# CÁLCULO PRINCIPAL
# ============================================================

def build_main_table(socios):
    acumulados = load_acumulados()
    audit = load_auditoria()

    df = socios.merge(
        acumulados,
        on="POS_CODE",
        how="left",
    )

    df["GENERADO"] = df["GENERADO"].fillna(0)

    # Detecta si el socio ya tuvo alguna carga auditada, incluso si
    # esa carga fue de monto 0.
    loaded_codes = set()
    if not audit.empty and "pos_code" in audit.columns:
        loaded_codes = {
            normalize_code(x)
            for x in audit["pos_code"].tolist()
            if normalize_code(x)
        }

    df["% CUMPLIMIENTO"] = (
        df["GENERADO"]
        / df["OBJETIVO"].replace(0, pd.NA)
        * 100
    ).fillna(0)

    df["FALTANTE"] = (
        df["OBJETIVO"]
        - df["GENERADO"]
    ).clip(lower=0)

    def state(row):
        code = normalize_code(row["POS_CODE"])

        if row["GENERADO"] == 0 and code not in loaded_codes:
            return "⚪ Sin datos"

        pct = row["% CUMPLIMIENTO"]

        if pct >= 80:
            return "🟢 Cumple / buen avance"
        if pct >= 40:
            return "🟡 En seguimiento"
        return "🔴 Bajo avance"

    df["ESTADO"] = df.apply(
        state,
        axis=1,
    )

    return df


def add_to_accumulated(confirmed):
    acumulados = load_acumulados()

    if acumulados.empty:
        acumulados = pd.DataFrame(
            columns=["POS_CODE", "GENERADO"]
        )

    grouped = (
        confirmed.groupby(
            "POS_CODE",
            as_index=False,
        )["monto"]
        .sum()
        .rename(columns={"monto": "AGREGAR"})
    )

    acumulados = acumulados.merge(
        grouped,
        on="POS_CODE",
        how="outer",
    )

    acumulados["GENERADO"] = (
        acumulados["GENERADO"].fillna(0)
        + acumulados["AGREGAR"].fillna(0)
    )

    acumulados = acumulados[
        ["POS_CODE", "GENERADO"]
    ]

    save_acumulados(acumulados)

    return grouped


def already_loaded_hash(file_hash_value):
    audit = load_auditoria()

    if audit.empty or "hash_archivo" not in audit.columns:
        return False

    return (
        audit["hash_archivo"].astype(str)
        == file_hash_value
    ).any()


def prepare_confirmed_dataframe(edited):
    rows = []

    for _, row in edited.iterrows():
        code = normalize_code(
            row.get("POS_CODE", "")
        )
        amount = safe_amount(
            row.get("AGREGAR", 0)
        )

        if code and amount != 0:
            rows.append(
                {
                    "POS_CODE": code,
                    "monto": amount,
                    "socio": str(
                        row.get("socio", "")
                    ),
                }
            )

    return pd.DataFrame(rows)


# ============================================================
# EXPORTACIÓN EXCEL
# ============================================================

def dataframe_to_excel_bytes(df):
    output = io.BytesIO()

    export = df.copy()

    export = export.rename(
        columns={
            "POS_OWNER": "Socio",
            "POS_CODE": "EH/POS",
            "CATEGORIA": "Categoría",
            "OBJETIVO": "Objetivo",
            "GENERADO": "Generado",
            "% CUMPLIMIENTO": "% Cumplimiento",
            "FALTANTE": "Faltante",
            "ESTADO": "Estado",
        }
    )

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        export.to_excel(
            writer,
            sheet_name="Seguimiento",
            index=False,
        )

        ws = writer.sheets["Seguimiento"]

        from openpyxl.styles import (
            Font,
            PatternFill,
            Alignment,
        )

        header_fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center"
            )

        # Semáforo de filas.
        for row in range(
            2,
            ws.max_row + 1,
        ):
            state_value = ws.cell(
                row=row,
                column=8,
            ).value

            if str(state_value).startswith("🟢"):
                fill = PatternFill(
                    "solid",
                    fgColor="C6EFCE",
                )
            elif str(state_value).startswith("🟡"):
                fill = PatternFill(
                    "solid",
                    fgColor="FFEB9C",
                )
            elif str(state_value).startswith("🔴"):
                fill = PatternFill(
                    "solid",
                    fgColor="FFC7CE",
                )
            else:
                fill = PatternFill(
                    "solid",
                    fgColor="E7E6E6",
                )

            for col in range(
                1,
                ws.max_column + 1,
            ):
                ws.cell(
                    row=row,
                    column=col,
                ).fill = fill

        widths = {
            "A": 40,
            "B": 14,
            "C": 20,
            "D": 12,
            "E": 12,
            "F": 18,
            "G": 12,
            "H": 28,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

    output.seek(0)
    return output.getvalue()


# ============================================================
# INICIO
# ============================================================

initialize_files()

st.title("🎯 Seguimiento de Objetivos de Socios")
st.caption(
    "Objetivos + generación acumulada + cumplimiento + auditoría de cargas"
)

socios = load_socios()

# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.header("⚙️ Configuración")

    usuario = st.text_input(
        "Usuario que realiza la carga",
        value=st.session_state.get(
            "usuario",
            "",
        ),
        placeholder="Ej. Vlady",
    )

    st.session_state["usuario"] = usuario

    st.divider()

    st.subheader("📥 Objetivos")

    objetivo_file = st.file_uploader(
        "Subir Objetivo_REPARADO.xlsx",
        type=["xlsx", "xls"],
        key="objetivo_file",
    )

    if objetivo_file is not None:

        objective_hash = file_hash(
            objetivo_file.getvalue()
        )

        if st.button(
            "Cargar / actualizar socios",
            type="secondary",
        ):
            try:
                new_socios = save_objectives_from_excel(
                    objetivo_file.getvalue()
                )

                st.success(
                    f"Se cargaron {len(new_socios)} socios."
                )

                st.rerun()

            except Exception as exc:
                st.error(str(exc))

    st.caption(
        "Si no cargas el Excel, se utilizan automáticamente "
        "los 14 socios de respaldo incluidos en la aplicación."
    )

    st.divider()

    st.subheader("💾 Persistencia local")

    st.caption(
        "La aplicación guarda acumulados y auditoría en CSV "
        "junto a app.py. No usa SQLite."
    )


# ============================================================
# PESTAÑAS
# ============================================================

tab_main, tab_file, tab_ocr, tab_audit = st.tabs(
    [
        "📊 Seguimiento",
        "📁 Vía A — Archivo",
        "📷 Vía B — Foto/OCR",
        "🧾 Auditoría",
    ]
)


# ============================================================
# TAB PRINCIPAL
# ============================================================

with tab_main:

    main_df = build_main_table(socios)

    total_obj = main_df["OBJETIVO"].sum()
    total_gen = main_df["GENERADO"].sum()
    total_faltante = main_df["FALTANTE"].sum()

    total_pct = (
        total_gen / total_obj * 100
        if total_obj
        else 0
    )

    c1, c2, c3, c4 = st.columns(4)

    c1.metric(
        "🎯 Objetivo total",
        f"{total_obj:,.0f}",
    )

    c2.metric(
        "📈 Generado",
        f"{total_gen:,.0f}",
    )

    c3.metric(
        "% Cumplimiento",
        f"{total_pct:.1f}%",
    )

    c4.metric(
        "⚠️ Faltante",
        f"{total_faltante:,.0f}",
    )

    st.divider()

    display = main_df[
        [
            "POS_OWNER",
            "POS_CODE",
            "CATEGORIA",
            "OBJETIVO",
            "GENERADO",
            "% CUMPLIMIENTO",
            "FALTANTE",
            "ESTADO",
        ]
    ].copy()

    display = display.rename(
        columns={
            "POS_OWNER": "Socio",
            "POS_CODE": "EH/POS",
            "CATEGORIA": "Categoría",
            "OBJETIVO": "Objetivo",
            "GENERADO": "Generado",
            "% CUMPLIMIENTO": "% Cumplimiento",
            "FALTANTE": "Faltante",
            "ESTADO": "Estado",
        }
    )

    total_row = pd.DataFrame([{
        "Socio": "TOTAL",
        "EH/POS": "",
        "Categoría": "",
        "Objetivo": total_obj,
        "Generado": total_gen,
        "% Cumplimiento": total_pct,
        "Faltante": total_faltante,
        "Estado": "📊 TOTAL GENERAL",
    }])

    display = pd.concat(
        [display, total_row],
        ignore_index=True,
    )

    st.dataframe(
        display,
        use_container_width=True,
        hide_index=True,
        column_config={
            "Objetivo": st.column_config.NumberColumn(
                format="%.0f"
            ),
            "Generado": st.column_config.NumberColumn(
                format="%.0f"
            ),
            "% Cumplimiento": st.column_config.NumberColumn(
                format="%.1f%%"
            ),
            "Faltante": st.column_config.NumberColumn(
                format="%.0f"
            ),
        },
    )

    excel_bytes = dataframe_to_excel_bytes(
        main_df
    )

    st.download_button(
        "📥 Exportar a Excel",
        data=excel_bytes,
        file_name=(
            "seguimiento_objetivos_socios.xlsx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


# ============================================================
# VÍA A
# ============================================================

with tab_file:

    st.subheader(
        "📁 Vía A — Subir archivo de ventas"
    )

    st.write(
        "El archivo debe contener una columna POS_CODE o EH "
        "y una columna de monto generado."
    )

    sales_file = st.file_uploader(
        "CSV o Excel de ventas",
        type=[
            "csv",
            "xlsx",
            "xls",
        ],
        key="sales_file",
    )

    if sales_file is not None:

        sales_bytes = sales_file.getvalue()
        sales_hash = file_hash(
            sales_bytes
        )

        if already_loaded_hash(
            sales_hash
        ):
            st.warning(
                "⚠️ Este archivo ya fue cargado anteriormente. "
                "No se sumará automáticamente."
            )

            duplicate_decision = st.radio(
                "¿Qué deseas hacer?",
                [
                    "Cancelar",
                    "Revisar y confirmar nuevamente",
                ],
                key="duplicate_file_decision",
            )

            duplicate_allowed = (
                duplicate_decision
                == "Revisar y confirmar nuevamente"
            )
        else:
            duplicate_allowed = True

        if duplicate_allowed:

            try:
                raw_sales = read_sales_file(
                    sales_bytes,
                    sales_file.name,
                )

                resolved, unidentified = resolve_rows(
                    raw_sales.to_dict("records"),
                    socios,
                )

                preview = build_preview_table(
                    resolved,
                    socios,
                )

                st.markdown(
                    "### Vista previa — antes de sumar"
                )

                if preview.empty:
                    st.error(
                        "No se identificó ninguna fila válida."
                    )
                else:
                    st.dataframe(
                        preview,
                        use_container_width=True,
                        hide_index=True,
                    )

                    st.info(
                        f"Se van a agregar "
                        f"**{preview['AGREGAR'].sum():,.0f}** "
                        f"en total."
                    )

                if not unidentified.empty:

                    st.warning(
                        "⚠️ Hay registros no identificados. "
                        "No se sumarán hasta que los corrijas."
                    )

                    st.dataframe(
                        unidentified,
                        use_container_width=True,
                        hide_index=True,
                    )

                    st.caption(
                        "Corrige esos registros manualmente "
                        "en el archivo y vuelve a cargarlo."
                    )

                if (
                    not preview.empty
                    and unidentified.empty
                ):

                    confirm = st.checkbox(
                        "Confirmo que esta carga es correcta "
                        "y deseo sumarla al acumulado.",
                        key="confirm_file_load",
                    )

                    if st.button(
                        "➕ Confirmar y sumar carga",
                        type="primary",
                        disabled=not confirm,
                    ):

                        grouped = add_to_accumulated(
                            resolved[
                                [
                                    "POS_CODE",
                                    "monto",
                                    "socio",
                                ]
                            ]
                        )

                        now = datetime.now().isoformat(
                            timespec="seconds"
                        )

                        audit_rows = []

                        for _, item in grouped.iterrows():

                            socio_name = (
                                preview.loc[
                                    preview["POS_CODE"]
                                    == item["POS_CODE"],
                                    "socio",
                                ].iloc[0]
                            )

                            audit_rows.append(
                                {
                                    "fecha_hora": now,
                                    "usuario": usuario
                                    or "No especificado",
                                    "archivo": sales_file.name,
                                    "tipo": "archivo",
                                    "hash_archivo": sales_hash,
                                    "pos_code": item["POS_CODE"],
                                    "socio": socio_name,
                                    "monto_sumado": item["AGREGAR"],
                                    "estado": "SUMADO",
                                }
                            )

                        append_auditoria(
                            audit_rows
                        )

                        st.success(
                            "✅ Carga sumada correctamente."
                        )

                        st.rerun()

            except Exception as exc:
                st.error(str(exc))


# ============================================================
# VÍA B — OCR
# ============================================================

with tab_ocr:

    st.subheader(
        "📷 Vía B — Foto / captura + OCR"
    )

    image_file = st.file_uploader(
        "Sube una foto/captura de ventas",
        type=[
            "png",
            "jpg",
            "jpeg",
            "webp",
        ],
        key="ocr_file",
    )

    if image_file is not None:

        image_bytes = image_file.getvalue()

        st.image(
            image_bytes,
            caption=image_file.name,
            use_container_width=True,
        )

        image_hash = file_hash(
            image_bytes
        )

        if already_loaded_hash(
            image_hash
        ):
            st.warning(
                "⚠️ Esta misma imagen ya fue cargada anteriormente. "
                "No se sumará automáticamente."
            )

        if st.button(
            "🔎 Ejecutar OCR",
            type="primary",
        ):

            rows, ocr_message = ocr_extract_rows(
                image_bytes
            )

            if not rows:
                st.warning(
                    ocr_message
                )

                manual_rows = pd.DataFrame(
                    [
                        {
                            "identificador": "",
                            "nombre_ocr": "",
                            "monto": 0,
                            "fuente": "manual",
                            "observacion": "",
                        }
                    ]
                )

            else:
                st.success(
                    f"OCR detectó {len(rows)} filas."
                )

                st.session_state[
                    "ocr_raw_rows"
                ] = rows

                st.session_state[
                    "ocr_message"
                ] = ocr_message

        if (
            "ocr_raw_rows" in st.session_state
        ):

            st.markdown(
                "### Revisar / corregir OCR"
            )

            ocr_df = pd.DataFrame(
                st.session_state[
                    "ocr_raw_rows"
                ]
            )

            edited_ocr = st.data_editor(
                ocr_df[
                    [
                        "identificador",
                        "nombre_ocr",
                        "monto",
                        "fuente",
                        "observacion",
                    ]
                ],
                num_rows="dynamic",
                use_container_width=True,
                hide_index=True,
                column_config={
                    "identificador": st.column_config.TextColumn(
                        "POS_CODE / EH"
                    ),
                    "nombre_ocr": st.column_config.TextColumn(
                        "Nombre leído"
                    ),
                    "monto": st.column_config.NumberColumn(
                        "Monto",
                        min_value=0,
                        step=1,
                    ),
                    "fuente": st.column_config.SelectboxColumn(
                        "Fuente",
                        options=[
                            "ocr",
                            "manual",
                        ],
                    ),
                    "observacion": st.column_config.TextColumn(
                        "Observación"
                    ),
                },
                key="ocr_editor",
            )

            resolved, unidentified = resolve_rows(
                edited_ocr.to_dict("records"),
                socios,
            )

            preview = build_preview_table(
                resolved,
                socios,
            )

            st.markdown(
                "### Vista previa — antes de sumar"
            )

            if not preview.empty:

                st.dataframe(
                    preview,
                    use_container_width=True,
                    hide_index=True,
                )

                st.info(
                    f"Total a agregar: "
                    f"**{preview['AGREGAR'].sum():,.0f}**"
                )

            if not unidentified.empty:

                st.warning(
                    "Hay filas que todavía no se pudieron "
                    "identificar por POS_CODE o nombre."
                )

                st.dataframe(
                    unidentified,
                    use_container_width=True,
                    hide_index=True,
                )

                st.caption(
                    "Corrige POS_CODE/EH o nombre en la tabla "
                    "superior y vuelve a revisar."
                )

            if (
                not preview.empty
                and unidentified.empty
            ):

                confirm_ocr = st.checkbox(
                    "Confirmo que la vista previa es correcta "
                    "y deseo sumar esta carga.",
                    key="confirm_ocr_load",
                )

                if st.button(
                    "➕ Confirmar y sumar OCR",
                    type="primary",
                    disabled=not confirm_ocr,
                ):

                    grouped = add_to_accumulated(
                        resolved[
                            [
                                "POS_CODE",
                                "monto",
                                "socio",
                            ]
                        ]
                    )

                    now = datetime.now().isoformat(
                        timespec="seconds"
                    )

                    audit_rows = []

                    for _, item in grouped.iterrows():

                        socio_match = preview[
                            preview["POS_CODE"]
                            == item["POS_CODE"]
                        ]

                        socio_name = (
                            socio_match["socio"].iloc[0]
                            if not socio_match.empty
                            else ""
                        )

                        audit_rows.append(
                            {
                                "fecha_hora": now,
                                "usuario": usuario
                                or "No especificado",
                                "archivo": image_file.name,
                                "tipo": "ocr",
                                "hash_archivo": image_hash,
                                "pos_code": item["POS_CODE"],
                                "socio": socio_name,
                                "monto_sumado": item["AGREGAR"],
                                "estado": "SUMADO",
                            }
                        )

                    append_auditoria(
                        audit_rows
                    )

                    st.success(
                        "✅ Carga OCR sumada correctamente."
                    )

                    st.session_state.pop(
                        "ocr_raw_rows",
                        None,
                    )

                    st.rerun()


# ============================================================
# AUDITORÍA
# ============================================================

with tab_audit:

    st.subheader(
        "🧾 Histórico / Auditoría de cargas"
    )

    audit = load_auditoria()

    if audit.empty:

        st.info(
            "Todavía no hay cargas registradas."
        )

    else:

        st.dataframe(
            audit.sort_values(
                "fecha_hora",
                ascending=False,
            ),
            use_container_width=True,
            hide_index=True,
        )

        st.download_button(
            "📥 Descargar auditoría CSV",
            data=audit.to_csv(
                index=False,
                encoding="utf-8-sig",
            ),
            file_name="auditoria_cargas.csv",
            mime="text/csv",
        )

    st.caption(
        "Cada registro conserva fecha/hora, usuario, archivo, "
        "hash del archivo, POS_CODE y monto sumado."
    )
